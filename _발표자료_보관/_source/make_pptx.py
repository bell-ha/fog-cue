#!/usr/bin/env python3
"""13장 SVG → 글자 편집이 되는 .pptx

PowerPoint의 SVG 도형 변환은 글자를 윤곽선으로 바꿔버려 편집이 안 된다.
그래서 이렇게 만든다:
  · 그래프 · 사진 · 도형  → 배경 PNG 한 장 (_bg/)
  · 글자 전부            → PowerPoint 텍스트 상자로 복원
글자는 클릭해서 바로 고칠 수 있고, 폰트·크기·색도 PowerPoint에서 바꾼다.
"""
import glob, html, os, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SVG_W, SVG_H = 1600, 900
W, H = Emu(12192000), Emu(6858000)
U = 12192000 / SVG_W                 # SVG 단위 → EMU
PT = 960 / SVG_W                     # SVG 단위 → pt (슬라이드 폭 960pt)

def attr(tag, name, default=None):
    m = re.search(rf'\b{name}="([^"]*)"', tag)
    return m.group(1) if m else default

def pick_font(family):
    f = (family or "").lower()
    if "mono" in f or "d2coding" in f: return "IBM Plex Mono"
    return "Pretendard"

def parse_texts(svg):
    """<text> 하나를 (줄 목록, 스타일)로 편다. tspan dy는 줄바꿈으로 본다."""
    out = []
    for m in re.finditer(r'<text\b([^>]*)>(.*?)</text>', svg, re.S):
        tag, inner = m.group(1), m.group(2)
        g = lambda k, d=None: attr(tag, k, d)
        x  = float(g("x", 0)); y = float(g("y", 0))
        size = float(g("font-size", 20))
        fill = g("fill", "#111111")
        weight = g("font-weight", "400")
        anchor = g("text-anchor", "start")
        family = pick_font(g("font-family", ""))
        ls = g("letter-spacing")

        # 첫 줄 = tspan 앞의 직접 텍스트
        head = re.split(r'<tspan\b', inner, 1)[0]
        lines = [(html.unescape(re.sub(r'<[^>]+>', '', head)).strip(), 0.0)]
        for t in re.finditer(r'<tspan\b([^>]*)>(.*?)</tspan>', inner, re.S):
            dy = float(attr(t.group(1), "dy", 0) or 0)
            tx = html.unescape(re.sub(r'<[^>]+>', '', t.group(2))).strip()
            # tspan이 색을 덮어쓰면 그 줄만 다른 색
            c = attr(t.group(1), "fill", None)
            lines.append((tx, dy, c))
        lines = [l for l in lines if l[0]]
        if not lines: continue
        out.append(dict(x=x, y=y, size=size, fill=fill, weight=weight,
                        anchor=anchor, family=family, ls=ls, lines=lines))
    return out

def rgb(c):
    c = (c or "#111111").lstrip("#")
    if len(c) == 3: c = "".join(ch*2 for ch in c)
    try: return RGBColor.from_string(c.upper())
    except Exception: return RGBColor.from_string("111111")

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]
total = 0

for f in sorted(glob.glob("[0-9][0-9].svg")):
    n = os.path.basename(f)[:2]
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(f"_bg/{n}.png", 0, 0, width=W, height=H)

    for t in parse_texts(open(f, encoding="utf-8").read()):
        size, x, y = t["size"], t["x"], t["y"]
        nlines = len(t["lines"])
        dys = [l[1] for l in t["lines"][1:]]
        span = sum(dys)
        # SVG y는 베이스라인. 텍스트 상자 상단 = 베이스라인 - 어센트
        top = (y - size * 0.86) * U
        left_pad = size * 0.14                    # 내부 여백 보정
        box_w = SVG_W * 0.98
        if t["anchor"] == "middle":
            left = (x - box_w/2) * U; align = PP_ALIGN.CENTER
        elif t["anchor"] == "end":
            left = (x - box_w) * U;   align = PP_ALIGN.RIGHT
        else:
            left = (x - left_pad) * U; align = PP_ALIGN.LEFT

        tb = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)),
                                      Emu(int(box_w*U)), Emu(int((size*1.35 + span)*U)))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP

        for i, ln in enumerate(t["lines"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if i > 0:
                p.line_spacing = Pt(ln[1] * PT)
            r = p.add_run(); r.text = ln[0]
            fnt = r.font
            fnt.name = t["family"]
            fnt.size = Pt(round(size * PT, 1))
            fnt.bold = t["weight"] in ("600", "700", "bold")
            col = ln[2] if len(ln) > 2 and ln[2] else t["fill"]
            fnt.color.rgb = rgb(col)
        total += 1

out = "동결_감지_자극_시스템.pptx"
prs.save(out)
print(f"  {out}")
print(f"  13장 · 텍스트 상자 {total}개 · {os.path.getsize(out)/1024/1024:.1f} MB")
